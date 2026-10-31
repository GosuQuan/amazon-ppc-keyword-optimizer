"""Build the white paper and the round archive from the markdown sources."""
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / 'docs'
ARCHIVE = DOCS / '迭代留痕'
ARCHIVE.mkdir(exist_ok=True)


def font(run, size=11, bold=False):
    run.font.name = 'Microsoft YaHei'
    run.font.size = Pt(size)
    run.bold = bold
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')


def add_markdown(document, text):
    for raw in text.splitlines():
        line = raw.rstrip()
        if not line:
            document.add_paragraph('')
            continue
        if line.startswith('# '):
            paragraph = document.add_paragraph()
            font(paragraph.add_run(line[2:]), 18, True)
        elif line.startswith('## '):
            paragraph = document.add_paragraph()
            font(paragraph.add_run(line[3:]), 14, True)
        elif line.startswith('- '):
            paragraph = document.add_paragraph(style='List Bullet')
            font(paragraph.add_run(line[2:]))
        else:
            paragraph = document.add_paragraph()
            font(paragraph.add_run(line))


def build(source, target):
    document = Document()
    add_markdown(document, source.read_text(encoding='utf-8'))
    document.save(target)
    reopened = Document(target)
    text = '\n'.join(paragraph.text for paragraph in reopened.paragraphs)
    if not any(marker in text for marker in ('亚马逊广告优化平台白皮书', '本轮技术思考', '平台需求拆分', '方案直达', '模型提示词', '可核对花费')):
        raise SystemExit('document missing expected heading: ' + target.name)
    print(target)


build(DOCS / '产品白皮书.md', DOCS / '产品白皮书.docx')
prompt = ARCHIVE / '2026-11-01-模型提示词指标.md'
prompt.write_text('# 2026-11-01 模型提示词指标\n\n' + (DOCS / '模型提示词指标测试边界.md').read_text(encoding='utf-8') + '\n这样每次扣积分的模型调用都带上已经算出的指标。未调用真实模型，也未验证生成词的广告效果。\n', encoding='utf-8')
build(prompt, ARCHIVE / '2026-11-01-模型提示词指标.docx')
claim = ARCHIVE / '2026-11-01-列表可核对花费.md'
claim.write_text('# 2026-11-01 列表可核对花费\n\n' + (DOCS / '列表可核对花费测试边界.md').read_text(encoding='utf-8'), encoding='utf-8')
build(claim, ARCHIVE / '2026-11-01-列表可核对花费.docx')
